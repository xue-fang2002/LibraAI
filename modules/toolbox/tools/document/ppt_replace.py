"""PPT 批量替换文字：对每个幻灯片里的文本框做查找替换。"""
import os

from core.response import ok, fail
from .._common import in_out, input_files

TOOL_ID = "document_ppt_replace"
CATEGORY = "document"
LABEL = "PPT 批量替换"
NEEDS_FILES = True
PARAMS = [
    {"name": "find_text", "label": "查找文字", "type": "text"},
    {"name": "replace_text", "label": "替换为", "type": "text"},
]


def _replace_in_shape(shape, old, new):
    count = 0
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if old in run.text:
                    count += run.text.count(old)
                    run.text = run.text.replace(old, new)
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        if old in run.text:
                            count += run.text.count(old)
                            run.text = run.text.replace(old, new)
    return count


def run(temp_id, params):
    try:
        import pptx
    except ImportError:
        return fail("服务器未安装 python-pptx，请联系管理员安装: pip install python-pptx", code=503)

    params = params or {}
    old = (params.get("find_text") or "")
    new = params.get("replace_text") or ""
    if old == "":
        return fail("请填写要查找的文字", code=400)

    in_dir, out_dir = in_out(temp_id)
    results = []
    for fpath in input_files(in_dir):
        if not fpath.lower().endswith(".pptx"):
            continue
        base = os.path.splitext(os.path.basename(fpath))[0]
        out_name = "replaced_" + base + ".pptx"
        out_path = os.path.join(out_dir, out_name)
        try:
            prs = pptx.Presentation(fpath)
            count = 0
            for slide in prs.slides:
                for shape in slide.shapes:
                    count += _replace_in_shape(shape, old, new)
            prs.save(out_path)
            results.append({"name": out_name, "size": os.path.getsize(out_path),
                            "replaced": count})
        except Exception as e:
            results.append({"name": os.path.basename(fpath), "error": str(e)})

    if not results:
        return fail("未找到可处理的 PPT 文件", code=400)
    return ok(data={"results": results})
