"""PPT 工具：提取文字 / 提取图片 / 转 PDF。

转 PDF 走 Office COM（_common.office_to_pdf），需 Windows + 已装 Office。
"""
import os

from core.response import ok, fail
from .._common import in_out, input_files, office_to_pdf

TOOL_ID = "document_ppt_tools"
CATEGORY = "document"
LABEL = "PPT 工具"
NEEDS_FILES = True
PARAMS = [
    {"name": "action", "label": "动作", "type": "enum",
     "options": ["extract_text", "extract_images", "to_pdf"],
     "default": "extract_text"},
]


def run(temp_id, params):
    params = params or {}
    action = params.get("action", "extract_text")
    try:
        import pptx
    except ImportError:
        return fail("服务器未安装 python-pptx，请联系管理员安装: pip install python-pptx", code=503)

    in_dir, out_dir = in_out(temp_id)
    files = [f for f in input_files(in_dir) if f.lower().endswith(".pptx")]

    if action == "to_pdf":
        return office_to_pdf(files, out_dir, "PowerPoint.Application", ".pdf")

    if not files:
        return fail("未找到 .pptx 文件", code=400)

    if action == "extract_images":
        results = []
        for fpath in files:
            base = os.path.splitext(os.path.basename(fpath))[0]
            try:
                prs = pptx.Presentation(fpath)
                idx = 0
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if shape.shape_type == 13 and hasattr(shape, "image"):
                            idx += 1
                            ext = shape.image.ext or "png"
                            out_name = f"{base}_img{idx}.{ext}"
                            out_path = os.path.join(out_dir, out_name)
                            with open(out_path, "wb") as fh:
                                fh.write(shape.image.blob)
                            results.append({"name": out_name,
                                            "size": os.path.getsize(out_path)})
                if idx == 0:
                    results.append({"name": os.path.basename(fpath),
                                    "error": "未找到内嵌图片"})
            except Exception as e:
                results.append({"name": os.path.basename(fpath), "error": str(e)})
        return ok(data={"results": results})

    lines = []
    for fpath in files:
        try:
            prs = pptx.Presentation(fpath)
            lines.append(f"=== {os.path.basename(fpath)} ===")
            for i, slide in enumerate(prs.slides):
                lines.append(f"--- 第 {i+1} 页 ---")
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        t = shape.text_frame.text.strip()
                        if t:
                            lines.append(t)
        except Exception as e:
            lines.append(f"{os.path.basename(fpath)} 提取失败: {e}")
    return ok(data={"result": "\n".join(lines) or "未提取到文字"})
