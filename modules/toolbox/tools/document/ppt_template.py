"""PPT 模板化批量生成：按数据表（Excel）每行替换占位符，生成多个 PPT。

占位符写成 {{字段名}}，与数据表列名一一对应。
"""
import os

from core.response import ok, fail
from .._common import in_out, input_files
from ._excel_common import read_table, list_excels

TOOL_ID = "document_ppt_template"
CATEGORY = "document"
LABEL = "PPT 模板生成"
NEEDS_FILES = True
PARAMS = [
    {"name": "name_field", "label": "用作文件名的列（留空=行号）", "type": "text"},
]


def _replace_shape(shape, mapping):
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                for k, v in mapping.items():
                    if k in run.text:
                        run.text = run.text.replace(k, str(v))


def _safe(name):
    return "".join(c for c in str(name) if c not in r'\\/:*?"<>|').strip() or "slide"


def run(temp_id, params):
    try:
        import pptx
    except ImportError:
        return fail("服务器未安装 python-pptx，请联系管理员安装: pip install python-pptx", code=503)

    params = params or {}
    in_dir, out_dir = in_out(temp_id)

    templates = [f for f in input_files(in_dir) if f.lower().endswith(".pptx")]
    datas = list_excels(in_dir)
    if not templates:
        return fail("请上传一个 .pptx 模板", code=400)
    if not datas:
        return fail("请上传一个数据表（xlsx/csv），第一行为字段名", code=400)

    tpl_path = templates[0]
    df = read_table(datas[0])
    name_field = (params.get("name_field") or "").strip()

    results = []
    for idx, (_, row) in enumerate(df.iterrows()):
        mapping = {"{{" + str(k) + "}}": ("" if v != v else str(v)) for k, v in row.items()}
        prs = pptx.Presentation(tpl_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                _replace_shape(shape, mapping)

        fname = str(row.get(name_field, idx + 1)) if name_field else str(idx + 1)
        out_name = _safe(fname) + ".pptx"
        out_path = os.path.join(out_dir, out_name)
        prs.save(out_path)
        results.append({"name": out_name, "size": os.path.getsize(out_path)})

    return ok(data={"results": results})
