"""Word 工具：批量替换 / 合并 / 转 PDF。

- replace：全量替换（段落 + 表格 + 页眉页脚，跨 run 合并后匹配，避免关键词被
  run 边界拆开而漏替换）。
- merge：合并多个 docx，逐 run 复制，尽量保留加粗/斜体/下划线/字体。
- to_pdf：走 Office COM / LibreOffice（_common.office_to_pdf）。
"""
import os

from core.response import ok, fail
from .._common import in_out, input_files, office_to_pdf

TOOL_ID = "document_word_tools"
CATEGORY = "document"
LABEL = "Word 工具"
NEEDS_FILES = True
PARAMS = [
    {"name": "action", "label": "动作", "type": "enum",
     "options": ["replace", "merge", "to_pdf"], "default": "replace"},
    {"name": "find_text", "label": "查找文字", "type": "text"},
    {"name": "replace_text", "label": "替换为", "type": "text"},
]


def _replace_in_paragraph(paragraph, old, new):
    """段落级跨 run 替换：合并 runs 文本后整体替换，写回第一个 run。"""
    full = "".join(r.text for r in paragraph.runs)
    if old not in full:
        return 0
    count = full.count(old)
    full = full.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = full
        for r in paragraph.runs[1:]:
            r.text = ""
    return count


def _replace_in_document(doc, old, new):
    total = 0
    for p in doc.paragraphs:
        total += _replace_in_paragraph(p, old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    total += _replace_in_paragraph(p, old, new)
    for section in doc.sections:
        for p in section.header.paragraphs:
            total += _replace_in_paragraph(p, old, new)
        for p in section.footer.paragraphs:
            total += _replace_in_paragraph(p, old, new)
    return total


def _copy_run(src_run, dst_run):
    dst_run.bold = src_run.bold
    dst_run.italic = src_run.italic
    dst_run.underline = src_run.underline
    if src_run.font.name:
        dst_run.font.name = src_run.font.name
    if src_run.font.size:
        dst_run.font.size = src_run.font.size
    if src_run.font.color and src_run.font.color.rgb:
        dst_run.font.color.rgb = src_run.font.color.rgb


def _merge_docx(base, sub):
    """把 sub 追加到 base，尽量保留 run 级格式（逐 run 复制）。"""
    base.add_page_break()
    for p in sub.paragraphs:
        new_p = base.add_paragraph(style=p.style)
        for r in p.runs:
            new_r = new_p.add_run(r.text)
            _copy_run(r, new_r)


def run(temp_id, params):
    params = params or {}
    action = params.get("action", "replace")
    try:
        import docx
    except ImportError:
        return fail("服务器未安装 python-docx，请联系管理员安装: pip install python-docx", code=503)

    in_dir, out_dir = in_out(temp_id)
    files = [f for f in input_files(in_dir) if f.lower().endswith(".docx")]

    if action == "to_pdf":
        return office_to_pdf(files, out_dir, "Word.Application", ".pdf")

    if not files:
        return fail("未找到 .docx 文件", code=400)

    if action == "merge":
        try:
            merged = docx.Document(files[0])
            for f in files[1:]:
                _merge_docx(merged, docx.Document(f))
            out_path = os.path.join(out_dir, "merged.docx")
            merged.save(out_path)
            return ok(data={"results": [{"name": "merged.docx",
                                         "size": os.path.getsize(out_path)}]})
        except Exception as e:
            return fail(f"合并失败：{e}", code=500)

    old = (params.get("find_text") or "").strip()
    new = params.get("replace_text") or ""
    if not old:
        return fail("请输入要查找的文字", code=400)

    results = []
    for fpath in files:
        fname = os.path.basename(fpath)
        try:
            doc = docx.Document(fpath)
            count = _replace_in_document(doc, old, new)
            out_name = "replaced_" + fname
            out_path = os.path.join(out_dir, out_name)
            doc.save(out_path)
            results.append({"name": out_name, "size": os.path.getsize(out_path),
                            "replaced": count})
        except Exception as e:
            results.append({"name": fname, "error": str(e)})
    return ok(data={"results": results})
