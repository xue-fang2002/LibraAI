"""文档格式转换聚合入口：一个工具覆盖所有「源格式 → 目标格式」组合。

为什么不拆成「PDF 转 Word」「Word 转 PDF」「HTML 转 PDF」……一堆小工具？
- 工具位是稀缺资源（每多一个，二级导航就长一截），而组合数会随格式数平方增长；
- 用户的心智模型是「我要把这个文件变成 X」，不是「我要用那个转换工具」。

因此这里只暴露「源格式 / 目标格式」两个下拉，内部按组合路由到具体实现。
不支持的组合明确报错，绝不静默返回空文件。
"""
import os

from core.response import ok, fail
from .._common import in_out, ordered_files

TOOL_ID = "document_convert"
CATEGORY = "document"
LABEL = "文档格式转换"
NEEDS_FILES = True
PARAMS = [
    {"name": "src", "label": "源格式（auto=按扩展名自动判断）", "type": "select"},
    {"name": "dst", "label": "目标格式", "type": "select"},
]

OFFICE_EXTS = (".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx")
TEXT_EXTS = (".txt", ".md", ".markdown", ".html", ".htm")
TABLE_EXTS = (".xlsx", ".xlsm", ".csv", ".tsv")
IMG_EXTS = (".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tiff")
DSTS = ("pdf", "docx", "xlsx", "csv", "txt", "md", "html", "pptx")


# ---------------------------------------------------------------- 目标：PDF
def _to_pdf(fpath, out_dir, ext):
    """任意可转换格式 → PDF。返回输出路径。"""
    base = os.path.splitext(os.path.basename(fpath))[0]
    out_path = os.path.join(out_dir, f"{base}.pdf")

    if ext in OFFICE_EXTS:
        from core.office_engine import convert_to_pdf, engine_status
        status = engine_status()
        if not status.get("engine"):
            raise RuntimeError(status.get("message") or "无可用文档转换引擎")
        return convert_to_pdf(fpath, out_dir)

    if ext in TEXT_EXTS:
        # 文本类先落成 docx（排版最简单可靠），再借道文档引擎转 PDF
        tmp_docx = os.path.join(out_dir, f"{base}__tmp.docx")
        _text_to_docx(fpath, tmp_docx, ext)
        from core.office_engine import convert_to_pdf, engine_status
        status = engine_status()
        if not status.get("engine"):
            os.path.exists(tmp_docx) and os.remove(tmp_docx)
            raise RuntimeError(status.get("message") or "无可用文档转换引擎")
        out = convert_to_pdf(tmp_docx, out_dir)
        os.path.exists(tmp_docx) and os.remove(tmp_docx)
        final = os.path.join(out_dir, f"{base}.pdf")
        if os.path.abspath(out) != os.path.abspath(final) and os.path.exists(out):
            os.replace(out, final)
        return final

    if ext in IMG_EXTS:
        from PIL import Image
        with Image.open(fpath) as im:
            pages = [im.convert("RGB")]
            pages[0].save(out_path, save_all=True)
        return out_path

    raise RuntimeError(f"暂不支持把 {ext} 转成 PDF")


# --------------------------------------------------------------- 目标：DOCX
def _text_to_docx(fpath, out_path, ext):
    """txt / md / html → docx（纯文本段落，不做复杂排版还原）。"""
    from docx import Document

    with open(fpath, "r", encoding="utf-8", errors="replace") as f:
        raw = f.read()
    if ext in (".html", ".htm"):
        import re
        raw = re.sub(r"<br\s*/?>", "\n", raw, flags=re.I)
        raw = re.sub(r"<[^>]+>", "", raw)
    doc = Document()
    for line in raw.splitlines():
        doc.add_paragraph(line)
    doc.save(out_path)
    return out_path


def _to_docx(fpath, out_dir, ext):
    base = os.path.splitext(os.path.basename(fpath))[0]
    out_path = os.path.join(out_dir, f"{base}.docx")

    if ext == ".pdf":
        import fitz
        from docx import Document
        doc = Document()
        with fitz.open(fpath) as pdf:
            for page in pdf:
                text = page.get_text("text") or ""
                for line in text.splitlines():
                    doc.add_paragraph(line)
                doc.add_page_break()
        doc.save(out_path)
        return out_path

    if ext in TEXT_EXTS:
        return _text_to_docx(fpath, out_path, ext)

    raise RuntimeError(f"暂不支持把 {ext} 转成 DOCX")


# ---------------------------------------------------------------- 目标：TXT
def _to_txt(fpath, out_dir, ext):
    base = os.path.splitext(os.path.basename(fpath))[0]
    out_path = os.path.join(out_dir, f"{base}.txt")

    if ext == ".pdf":
        import fitz
        buf = []
        with fitz.open(fpath) as pdf:
            for page in pdf:
                buf.append(page.get_text("text") or "")
        text = "\n".join(buf)
    elif ext == ".docx":
        from docx import Document
        doc = Document(fpath)
        text = "\n".join(p.text for p in doc.paragraphs)
    elif ext in TABLE_EXTS:
        from ._excel_common import read_table
        df = read_table(fpath)
        text = df.to_csv(index=False, sep="\t")
    elif ext in TEXT_EXTS:
        with open(fpath, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
    else:
        raise RuntimeError(f"暂不支持把 {ext} 转成 TXT")

    with open(out_path, "w", encoding="utf-8") as f:
        f.write(text)
    return out_path


# ------------------------------------------------------------- 目标：MD/HTML
def _to_md(fpath, out_dir, ext):
    base = os.path.splitext(os.path.basename(fpath))[0]
    out_path = os.path.join(out_dir, f"{base}.md")

    if ext == ".docx":
        from docx import Document
        doc = Document(fpath)
        lines = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or "").lower()
            if "heading 1" in style or style == "title":
                lines.append(f"# {t}")
            elif "heading 2" in style:
                lines.append(f"## {t}")
            elif "heading" in style:
                lines.append(f"### {t}")
            else:
                lines.append(t)
        text = "\n\n".join(lines)
    elif ext in TEXT_EXTS:
        with open(fpath, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
    else:
        raise RuntimeError(f"暂不支持把 {ext} 转成 Markdown")

    with open(out_path, "w", encoding="utf-8") as f:
        f.write(text)
    return out_path


def _to_html(fpath, out_dir, ext):
    base = os.path.splitext(os.path.basename(fpath))[0]
    out_path = os.path.join(out_dir, f"{base}.html")

    with open(fpath, "r", encoding="utf-8", errors="replace") as f:
        raw = f.read() if ext in TEXT_EXTS else ""

    if ext in (".md", ".markdown"):
        try:
            import markdown as md_lib
            html = md_lib.markdown(raw, extensions=["tables", "fenced_code"])
        except ImportError:
            import html as html_lib
            body = html_lib.escape(raw)
            html = f"<pre>{body}</pre>"
    elif ext in (".txt",):
        import html as html_lib
        html = f"<pre>{html_lib.escape(raw)}</pre>"
    elif ext == ".docx":
        from docx import Document
        import html as html_lib
        doc = Document(fpath)
        parts = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            tag = "h2" if "heading 1" in (p.style.name or "").lower() else "p"
            parts.append(f"<{tag}>{html_lib.escape(t)}</{tag}>")
        html = "\n".join(parts)
    else:
        raise RuntimeError(f"暂不支持把 {ext} 转成 HTML")

    if not html.lstrip().lower().startswith(("<!doctype", "<html")):
        html = (f"<!DOCTYPE html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\">"
                f"<title>{base}</title></head><body>{html}</body></html>")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)
    return out_path


# ------------------------------------------------------------ 目标：表格/PPT
def _to_table(fpath, out_dir, ext, dst):
    from ._excel_common import read_table, write_table
    base = os.path.splitext(os.path.basename(fpath))[0]
    if ext not in TABLE_EXTS:
        raise RuntimeError(f"暂不支持把 {ext} 转成 {dst}（目前仅表格格式互转）")
    out_path = os.path.join(out_dir, f"{base}.{dst}")
    write_table(read_table(fpath), out_path)
    return out_path


def _to_pptx(fpath, out_dir, ext):
    """目前只支持 PDF → PPTX（逐页渲染成图），复用 pdf_to_ppt 的实现。"""
    if ext != ".pdf":
        raise RuntimeError(f"暂不支持把 {ext} 转成 PPTX")
    from .pdf_to_ppt import convert_one
    return convert_one(fpath, out_dir, dpi=150)


def run(temp_id, params):
    params = params or {}
    dst = (params.get("dst") or "").strip().lower()
    src = (params.get("src") or "auto").strip().lower()
    if dst not in DSTS:
        return fail(f"不支持的目标格式: {dst}", code=400)

    in_dir, out_dir = in_out(temp_id)
    files = ordered_files(in_dir)
    if not files:
        return fail("请先上传要转换的文件", code=400)

    results = []
    for fpath in files:
        base = os.path.splitext(os.path.basename(fpath))[0]
        ext = os.path.splitext(fpath)[1].lower() if src == "auto" else f".{src.lstrip('.')}"
        if not ext:
            results.append({"name": base, "error": "无法判断源格式，请手动选择"})
            continue
        if ext.lstrip(".") == dst:
            results.append({"name": base, "error": f"源格式与目标格式相同（{dst}），无需转换"})
            continue
        try:
            if dst == "pdf":
                out = _to_pdf(fpath, out_dir, ext)
            elif dst == "docx":
                out = _to_docx(fpath, out_dir, ext)
            elif dst == "txt":
                out = _to_txt(fpath, out_dir, ext)
            elif dst == "md":
                out = _to_md(fpath, out_dir, ext)
            elif dst == "html":
                out = _to_html(fpath, out_dir, ext)
            elif dst in ("xlsx", "csv"):
                out = _to_table(fpath, out_dir, ext, dst)
            else:  # pptx
                out = _to_pptx(fpath, out_dir, ext)
            results.append({"name": os.path.basename(out), "size": os.path.getsize(out)})
        except Exception as e:
            results.append({"name": f"{base}.{dst}", "error": str(e)})

    return ok(data={"results": results})
