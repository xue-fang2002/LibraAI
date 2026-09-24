"""PDF 转 Word：提取每页文本，重组为 .docx（文本型，不保留复杂排版）。"""
import os

from core.response import ok, fail
from .._common import in_out, input_files

TOOL_ID = "document_pdf_to_word"
CATEGORY = "document"
LABEL = "PDF 转 Word"
NEEDS_FILES = True
PARAMS = []


def run(temp_id, params):
    try:
        import fitz
    except ImportError:
        return fail("服务器未安装 PyMuPDF，请联系管理员安装: pip install PyMuPDF", code=503)
    try:
        import docx
    except ImportError:
        return fail("服务器未安装 python-docx，请联系管理员安装: pip install python-docx", code=503)

    in_dir, out_dir = in_out(temp_id)
    results = []
    for fpath in input_files(in_dir):
        if not fpath.lower().endswith(".pdf"):
            continue
        base = os.path.splitext(os.path.basename(fpath))[0]
        out_name = base + ".docx"
        out_path = os.path.join(out_dir, out_name)
        try:
            pdf = fitz.open(fpath)
            document = docx.Document()
            document.add_heading(base, level=0)
            for i, page in enumerate(pdf):
                text = page.get_text("text").strip()
                if text:
                    document.add_heading(f"第 {i+1} 页", level=1)
                    for para in text.split("\n"):
                        para = para.strip()
                        if para:
                            document.add_paragraph(para)
            pdf.close()
            document.save(out_path)
            results.append({"name": out_name, "size": os.path.getsize(out_path)})
        except Exception as e:
            results.append({"name": os.path.basename(fpath), "error": str(e)})

    if not results:
        return fail("未找到可转换的 PDF 文件", code=400)
    return ok(data={"results": results})
