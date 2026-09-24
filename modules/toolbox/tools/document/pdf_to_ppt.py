"""PDF → PPT：把每页渲染成图片插入幻灯片，页码顺序不变。

排版不还原（不是矢量转换），定位是「快速把 PDF 内容搬进 PPT 汇报稿」。
"""
import os

from core.response import ok, fail
from .._common import in_out, input_files

TOOL_ID = "document_pdf_to_ppt"
CATEGORY = "document"
LABEL = "PDF 转 PPT"
NEEDS_FILES = True
PARAMS = [
    {"name": "dpi", "label": "渲染分辨率", "type": "select"},
]


def convert_one(fpath, out_dir, dpi=150):
    """把单个 PDF 转成 pptx，返回输出路径（供本工具与 document_convert 共用）。"""
    import fitz
    from pptx import Presentation
    from pptx.util import Emu

    base = os.path.splitext(os.path.basename(fpath))[0]
    out_path = os.path.join(out_dir, f"{base}.pptx")
    tmp_imgs = []
    try:
        prs = Presentation()
        blank = prs.slide_layouts[6]
        slide_w, slide_h = prs.slide_width, prs.slide_height

        with fitz.open(fpath) as doc:
            for i, page in enumerate(doc):
                pix = page.get_pixmap(dpi=dpi)
                img_path = os.path.join(out_dir, f"_ppt_{base}_{i}.png")
                pix.save(img_path)
                tmp_imgs.append(img_path)

                # 等比缩放居中（contain），不变形
                ratio = min(slide_w / pix.width, slide_h / pix.height)
                w, h = int(pix.width * ratio), int(pix.height * ratio)
                slide = prs.slides.add_slide(blank)
                slide.shapes.add_picture(
                    img_path,
                    Emu(int((slide_w - w) / 2)),
                    Emu(int((slide_h - h) / 2)),
                    Emu(w), Emu(h),
                )

        prs.save(out_path)
        return out_path
    finally:
        for p in tmp_imgs:
            try:
                os.remove(p)
            except OSError:
                pass


def run(temp_id, params):
    try:
        import fitz  # noqa: F401
    except ImportError:
        return fail("服务器未安装 PyMuPDF：pip install PyMuPDF", code=503)
    try:
        from pptx import Presentation  # noqa: F401
    except ImportError:
        return fail("服务器未安装 python-pptx：pip install python-pptx", code=503)

    params = params or {}
    try:
        dpi = int(params.get("dpi") or 150)
    except (TypeError, ValueError):
        dpi = 150
    dpi = max(72, min(dpi, 300))

    in_dir, out_dir = in_out(temp_id)
    files = [f for f in input_files(in_dir) if f.lower().endswith(".pdf")]
    if not files:
        return fail("未找到 PDF 文件", code=400)

    results = []
    for fpath in files:
        base = os.path.splitext(os.path.basename(fpath))[0]
        try:
            out_path = convert_one(fpath, out_dir, dpi)
            results.append({"name": os.path.basename(out_path),
                            "size": os.path.getsize(out_path)})
        except Exception as e:
            results.append({"name": f"{base}.pptx", "error": str(e)})

    return ok(data={"results": results})
