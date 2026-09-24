"""PPT 合并：把多个 .pptx 合并为一个（保留文字与形状，内嵌图片可能丢失）。"""
import copy
import os

from core.response import ok, fail
from .._common import in_out, input_files

TOOL_ID = "document_ppt_merge"
CATEGORY = "document"
LABEL = "PPT 合并"
NEEDS_FILES = True
PARAMS = []


def _copy_slide(dest_prs, src_slide):
    # 使用目标文件自身的布局（避免把源文件的 slideLayout 塞进目标包，
    # 否则保存时会出现重复的 layout part 导致文件损坏）
    layout = dest_prs.slide_layouts[0]
    dest_slide = dest_prs.slides.add_slide(layout)
    # 删除新 slide 自带的占位符
    for shp in list(dest_slide.shapes):
        shp._element.getparent().remove(shp._element)
    # 深拷贝来源 slide 的所有形状
    for shp in src_slide.shapes:
        dest_slide.shapes._spTree.append(copy.deepcopy(shp._element))
    return dest_slide


def run(temp_id, params):
    try:
        import pptx
    except ImportError:
        return fail("服务器未安装 python-pptx，请联系管理员安装: pip install python-pptx", code=503)

    in_dir, out_dir = in_out(temp_id)
    files = sorted(f for f in input_files(in_dir) if f.lower().endswith(".pptx"))
    if not files:
        return fail("未找到可合并的 PPT 文件", code=400)

    try:
        merged = pptx.Presentation(files[0])
        for fpath in files[1:]:
            prs = pptx.Presentation(fpath)
            for slide in prs.slides:
                _copy_slide(merged, slide)
        out_path = os.path.join(out_dir, "merged.pptx")
        merged.save(out_path)
    except Exception as e:
        return fail("合并失败：" + str(e), code=500)

    return ok(data={"results": [{"name": "merged.pptx", "size": os.path.getsize(out_path)}]})
