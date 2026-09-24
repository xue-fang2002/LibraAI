"""
办公工具 - Flask 路由层
重构说明（支持远程/其他电脑使用）：
  原版所有操作都依赖"用户在文本框填写服务器上的绝对路径"，因此只有部署代码的
  服务器本机才可用。改为"上传 zip（你的文件夹）→ 服务端在临时工作目录处理 →
  把结果目录打成 zip 下载回去"的统一模型，任何联网电脑都能真正使用。
  service 层函数完全复用（它们本来就接收目录/文件路径）。
"""
import os
import io
import csv
import uuid
import time
import shutil
import tempfile
import zipfile
from flask import render_template, request, send_file, jsonify

from modules.office_tools import bp
from modules.office_tools.service import (
    batch_create_folders, batch_create_files, batch_create_sheets,
    extract_items, extract_detailed_data,
    batch_rename_simple, batch_rename_by_excel,
    replace_cell_content, batch_excel_to_pdf, delete_sheets,
    batch_move_copy, extract_file_paths,
    generate_template, HAS_WIN32COM, EXCEL_EXTENSIONS,
)
from core.response import ok, fail
from core.exceptions import module_exception_guard
from core.auth import login_required
from core.office_engine import engine_status
from common.utils import is_safe_path, _safe_dirname
from . import ai_planner

UPLOAD_FOLDER = os.path.join(tempfile.gettempdir(), "office_tools_uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# 提取/路径类结果缓存：前端拿到 token，导出时回查（避免把大量条目再传一遍）
_RESULT_CACHE = {}
# 缓存与临时文件保留时长（秒），到期自动清理，避免长时运行内存/磁盘泄漏
_RESULT_TTL = 3600


def _purge_expired_cache():
    """清理过期的提取结果缓存及其临时工作目录，并回收上传目录下过期残留。"""
    now = time.time()
    expired = [k for k, v in _RESULT_CACHE.items() if now - v.get("ts", 0) > _RESULT_TTL]
    for k in expired:
        v = _RESULT_CACHE.pop(k, None)
        rd = (v or {}).get("root_dir")
        if rd and os.path.isdir(rd):
            try:
                shutil.rmtree(rd)
            except Exception:
                pass
    try:
        for fn in os.listdir(UPLOAD_FOLDER):
            fp = os.path.join(UPLOAD_FOLDER, fn)
            try:
                if now - os.path.getmtime(fp) > _RESULT_TTL:
                    if os.path.isfile(fp):
                        os.remove(fp)
                    elif os.path.isdir(fp):
                        shutil.rmtree(fp)
            except Exception:
                pass
    except Exception:
        pass


@bp.before_request
def _office_auto_cleanup():
    _purge_expired_cache()


# ======================== 通用辅助 ========================

def save_upload_file(file_obj):
    """保存上传文件到临时目录，返回完整路径"""
    ext = os.path.splitext(file_obj.filename)[1]
    filename = f"{uuid.uuid4().hex}{ext}"
    filepath = os.path.join(UPLOAD_FOLDER, filename)
    file_obj.save(filepath)
    return filepath


def _new_workdir():
    """创建一个本次操作专属的临时工作目录"""
    d = os.path.join(UPLOAD_FOLDER, uuid.uuid4().hex)
    os.makedirs(d, exist_ok=True)
    return d


def _require_zip(key="folder_zip"):
    f = request.files.get(key)
    if not f or not f.filename:
        raise ValueError("请上传文件夹压缩包（zip）")
    return f


def _decode_zip_name(name, flag_bits=0):
    """正确还原 zip 条目名（修复中文乱码）。

    Python 的 zipfile 对**未设置 UTF-8 标志位（bit 11）**的条目，默认按 cp437
    解码文件名。而 Windows 资源管理器 / 7-Zip 打的中文压缩包，文件名其实是 GBK
    字节且没有该标志位 —— 直接拿 namelist()/extractall 会得到 «─Ω╘┬» 这类乱码。
    这里：有 UTF-8 标志就用原串；否则把 cp437 串按 GBK 还原（gb18030 兜底更全）。
    """
    if not name:
        return name
    if flag_bits & 0x800:
        return name  # 已是 UTF-8，zipfile 已正确解码
    try:
        return name.encode("cp437").decode("gb18030")
    except Exception:
        return name


# generate_template() 实际支持的模板类型（service.py 的分支全集）。
# 新增类型时必须同步登记，否则 /api/template/<type> 会返回 400。
_TEMPLATE_TYPES = {"folder", "file", "sheet", "rename", "mapping"}


def _zip_target_is_safe(decoded_name, root):
    """判断解压目标是否落在 root 内（防 zip-slip）。

原先只校验 info.filename（cp437 原串），而实际落盘用的是
    _decode_zip_name() 还原出的 gb18030 串 —— 两者并不总是相同，精心构造的条目
    可以做到「校验的原串合法、落盘的还原串穿越」。这里改为校验**真正会落盘的串**，
    并对最终绝对路径做 commonpath 实证复核（顺带挡住盘符与 UNC）。
    """
    n = str(decoded_name or "").replace("\\", "/")
    if not n:
        return False
    if n.startswith("/") or n.startswith("//"):
        return False
    if len(n) > 1 and n[1] == ":":
        return False
    parts = []
    for seg in n.split("/"):
        if seg in ("", "."):
            continue
        if seg == "..":
            return False
        parts.append(seg)
    if not parts:
        return False
    try:
        target = os.path.abspath(os.path.join(root, *parts))
        return os.path.commonpath([root, target]) == root
    except ValueError:
        return False


def _extract_zip_to(f, work_dir):
    """把上传的 zip 解压到 work_dir（防路径穿越 + 中文文件名正确还原）"""
    data = f.read()
    root = os.path.abspath(work_dir)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        # 先按「实际会落盘的名字」全量校验，任何一个不合法就整包拒绝
        for info in z.infolist():
            decoded = _decode_zip_name(info.filename, info.flag_bits)
            if not _zip_target_is_safe(decoded, root):
                raise ValueError("压缩包含非法路径，已拒绝解压")
        # 逐条目解压，文件名按中文正确解码（extractall 会沿用 cp437 乱码名）
        for info in z.infolist():
            decoded = _decode_zip_name(info.filename, info.flag_bits)
            # 目录条目：只需建目录
            if decoded.endswith("/") or decoded.endswith("\\"):
                os.makedirs(os.path.join(work_dir, decoded), exist_ok=True)
                continue
            target = os.path.abspath(os.path.join(work_dir, decoded))
            # 落盘前二次实证复核
            if os.path.commonpath([root, target]) != root:
                raise ValueError("压缩包含非法路径，已拒绝解压")
            os.makedirs(os.path.dirname(target), exist_ok=True)
            with z.open(info) as src, open(target, "wb") as dst:
                dst.write(src.read())
    return work_dir


def _zip_workdir(work_dir, hint):
    """把工作目录打成 zip 放到 UPLOAD_FOLDER，返回文件名"""
    zip_name = f"{uuid.uuid4().hex}_{hint}.zip"
    zip_path = os.path.join(UPLOAD_FOLDER, zip_name)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for root, _, files in os.walk(work_dir):
            for fn in files:
                fp = os.path.join(root, fn)
                arc = os.path.relpath(fp, work_dir)
                z.write(fp, arc)
    return zip_name


def _download_url(zip_name):
    return f"/office_tools/api/download?file={zip_name}"


def _clean_logs(result, root):
    """把日志里的服务器临时绝对路径替换为相对显示，避免泄露服务端路径"""
    for e in result.get("log", []):
        if isinstance(e, dict) and "msg" in e:
            e["msg"] = e["msg"].replace(root, "").replace("\\", "/")
    return result


def _build_table_file(headers, rows, file_type):
    """根据表头+数据生成 csv/xlsx，返回下载文件名（在 UPLOAD_FOLDER）"""
    if file_type == "csv":
        filepath = os.path.join(UPLOAD_FOLDER, f"{uuid.uuid4().hex}.csv")
        with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            writer.writerows(rows)
    else:
        from openpyxl import Workbook
        filepath = os.path.join(UPLOAD_FOLDER, f"{uuid.uuid4().hex}.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for row in rows:
            ws.append(row)
        wb.save(filepath)
        wb.close()
    return os.path.basename(filepath)


# ======================== 页面路由 ========================

@bp.route("/")
@module_exception_guard("office_tools")
@login_required
def index():
    # 不再显式传 current_module_id：本模块已并入「工具箱 Hub」，归属导航项由
    # PluginScanner.HUB_PARENT_MODULE 在 app.py 统一归一化为 toolbox_hub。
    # （显式传参优先级高于全局注入，会把归一化结果覆盖掉，导致侧栏无任何高亮）
    return render_template("office_tools/index.html")


# ======================== AI 助手（计划预览 + 确认后执行） ========================

@bp.route("/api/ai/plan", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_ai_plan():
    """
    生成办公自动化执行计划。
    body: {"text": "需求描述", "caps": [{id,label,desc,params:[...]}, ...]}
    返回 plan: {tool, label, params, reply}（文件类输入不参与，仍由用户上传）
    """
    data = request.get_json(silent=True) or {}
    text = (data.get("text") or "").strip()
    caps = data.get("caps") or []
    if not text:
        return fail("请输入需求描述", code=400)
    if not caps:
        return fail("缺少能力清单", code=400)

    caps_by_id = {c.get("id"): c for c in caps if isinstance(c, dict) and c.get("id")}
    if not caps_by_id:
        return fail("能力清单为空", code=400)

    from modules.ai_center.internal.qa import llm_generate
    prompt = ai_planner.build_plan_prompt(text, list(caps_by_id.values()))
    raw = llm_generate(prompt, max_tokens=600, temperature=0.2)
    if not raw:
        return fail("模型未响应，请检查 AI 中心后端配置")

    plan = ai_planner.parse_plan(raw)
    clean, err = ai_planner.validate_plan(plan, caps_by_id)
    if err:
        return fail(f"计划解析失败：{err}", data={"raw": raw[:400]})
    if clean is None:
        reply = str((plan or {}).get("reply") or "没有找到匹配的功能")
        return ok(data={"plan": None, "reply": reply, "raw": raw[:400]})
    return ok(data={"plan": clean, "reply": clean.get("reply") or "", "raw": raw[:400]})


# ======================== 模块1：批量创建 ========================

def _safe_sub_dirs(work_dir, dirs_raw):
    """把用户填写的目标子目录清洗为 work_dir 内的安全路径（防路径穿越）。

    - 拒绝绝对路径 / 盘符开头 / 含 .. 或 . 分量的条目（直接报错，不静默跳过）
    - 每个分量过 _safe_dirname 清洗（剔除 \\/:*?\"<>|）
    - 最终 join 后强制 is_safe_path 校验，必须落在 work_dir 内
    """
    sub_dirs = [d.strip() for d in (dirs_raw or "").split("\n") if d.strip()]
    target_dirs = []
    for d in sub_dirs:
        norm = d.replace("/", "\\")
        parts = [p for p in norm.split("\\") if p]
        if not parts:
            continue
        if os.path.isabs(d) or (len(parts[0]) >= 2 and parts[0][1] == ":"):
            raise ValueError(f"目标子目录不允许绝对路径: {d}")
        if any(p in (".", "..") for p in parts):
            raise ValueError(f"目标子目录含非法路径分量: {d}")
        safe_parts = [_safe_dirname(p) for p in parts]
        full = os.path.normpath(os.path.join(work_dir, *safe_parts))
        if not is_safe_path(full, work_dir):
            raise ValueError(f"目标子目录越界: {d}")
        target_dirs.append(full)
    return target_dirs or [work_dir]


@bp.route("/api/mod1/folders", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_mod1_folders():
    """批量建文件夹：名称Excel + 目标子目录（相对工作根，留空=根）"""
    try:
        if "excel" not in request.files:
            return fail("请上传名称Excel文件")
        excel_file = request.files["excel"]
        column_header = request.form.get("column_header", "文件夹名称").strip()
        dirs_raw = (request.form.get("target_dirs") or "").strip()
        excel_path = save_upload_file(excel_file)
        work_dir = _new_workdir()
        target_dirs = _safe_sub_dirs(work_dir, dirs_raw)
        result = batch_create_folders(excel_path, column_header, target_dirs)
        os.remove(excel_path)
        zip_name = _zip_workdir(work_dir, "folders")
        _clean_logs(result, work_dir)
        return ok(data={**result, "download_url": _download_url(zip_name)})
    except Exception as e:
        return fail(str(e))


@bp.route("/api/mod1/files", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_mod1_files():
    """批量建文件：名称Excel + 目标子目录 + 文件类型"""
    try:
        if "excel" not in request.files:
            return fail("请上传名称Excel文件")
        excel_file = request.files["excel"]
        column_header = request.form.get("column_header", "文件名称").strip()
        dirs_raw = (request.form.get("target_dirs") or "").strip()
        file_type = request.form.get("file_type", "excel")
        excel_path = save_upload_file(excel_file)
        work_dir = _new_workdir()
        target_dirs = _safe_sub_dirs(work_dir, dirs_raw)
        result = batch_create_files(excel_path, column_header, target_dirs, file_type)
        os.remove(excel_path)
        zip_name = _zip_workdir(work_dir, "files")
        _clean_logs(result, work_dir)
        return ok(data={**result, "download_url": _download_url(zip_name)})
    except Exception as e:
        return fail(str(e))


@bp.route("/api/mod1/sheets", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_mod1_sheets():
    """批量建Excel分表：分表名称Excel + 需要加表的Excel文件（可多选，上传） + 可选模板"""
    try:
        if "excel" not in request.files:
            return fail("请上传分表名称Excel文件")
        excel_file = request.files["excel"]
        column_header = request.form.get("column_header", "分表名称").strip()
        fill = request.form.get("fill", "false").lower() == "true"
        targets = request.files.getlist("target_files")
        if not targets:
            return fail("请上传需要添加分表的Excel文件（可多选）")
        excel_path = save_upload_file(excel_file)
        work_dir = _new_workdir()
        target_paths = []
        for t in targets:
            tp = save_upload_file(t)
            dest = os.path.join(work_dir, os.path.basename(tp))
            os.replace(tp, dest)
            target_paths.append(dest)
        template_path = None
        if fill and "template" in request.files:
            template_file = request.files["template"]
            tpath = save_upload_file(template_file)
            tdest = os.path.join(work_dir, os.path.basename(tpath))
            os.replace(tpath, tdest)
            template_path = tdest
        result = batch_create_sheets(excel_path, column_header, target_paths, fill, template_path)
        os.remove(excel_path)
        zip_name = _zip_workdir(work_dir, "sheets")
        _clean_logs(result, work_dir)
        return ok(data={**result, "download_url": _download_url(zip_name)})
    except Exception as e:
        return fail(str(e))


# ======================== 模块2：提取清单 ========================

@bp.route("/api/mod2/extract", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_mod2_extract():
    """提取文件/文件夹清单：上传文件夹 zip"""
    try:
        folder_zip = _require_zip()
        mode = (request.form.get("mode") or "files").strip()
        recursive = request.form.get("recursive", "false").lower() == "true"
        suffix = (request.form.get("suffix") or "").strip()
        remove_suffix = request.form.get("remove_suffix", "false").lower() == "true"
        work_dir = _new_workdir()
        _extract_zip_to(folder_zip, work_dir)
        result = extract_items(work_dir, mode, recursive, suffix, remove_suffix)
        token = uuid.uuid4().hex
        _RESULT_CACHE[token] = {
            "items": result["items"],
            "root_dir": work_dir,
            "remove_suffix": remove_suffix,
            "ts": time.time(),
        }
        return ok(data={
            "items": result["items"],
            "names": result["names"],
            "count": result["count"],
            "token": token,
        })
    except Exception as e:
        return fail(str(e))


@bp.route("/api/mod2/export-detailed", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_mod2_export_detailed():
    """导出详细清单：用 token 回查提取结果"""
    try:
        token = request.form.get("token") or (request.get_json(silent=True) or {}).get("token")
        cached = _RESULT_CACHE.get(token)
        if not cached:
            return fail("提取结果已失效，请重新执行提取")
        items = cached["items"]
        root_dir = cached["root_dir"]
        remove_suffix = cached.get("remove_suffix", False)
        file_type = request.form.get("file_type") or (request.get_json(silent=True) or {}).get("file_type") or "xlsx"
        if not items:
            return fail("无提取内容")
        headers, rows = extract_detailed_data(items, root_dir, remove_suffix)
        fname = _build_table_file(headers, rows, file_type)
        return ok(data={"download_url": f"/office_tools/api/download?file={fname}"})
    except Exception as e:
        return fail(str(e))


# ======================== 模块3：批量改名 ========================

@bp.route("/api/mod3/rename-simple", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_mod3_rename_simple():
    """简单文件名替换：上传文件夹 zip"""
    try:
        folder_zip = _require_zip()
        old_text = (request.form.get("old_text") or "").strip()
        new_text = request.form.get("new_text") or ""
        if not old_text:
            return fail("查找文字不能为空")
        work_dir = _new_workdir()
        _extract_zip_to(folder_zip, work_dir)
        result = batch_rename_simple(work_dir, old_text, new_text)
        zip_name = _zip_workdir(work_dir, "renamed")
        _clean_logs(result, work_dir)
        return ok(data={**result, "download_url": _download_url(zip_name)})
    except Exception as e:
        return fail(str(e))


@bp.route("/api/mod3/rename-excel", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_mod3_rename_excel():
    """Excel对照表批量重命名：上传文件夹 zip + 对照Excel"""
    try:
        folder_zip = _require_zip()
        if "excel" not in request.files:
            return fail("请上传对照Excel文件")
        old_col = request.form.get("old_col", "原文件名").strip()
        new_col = request.form.get("new_col", "新文件名").strip()
        excel_path = save_upload_file(request.files["excel"])
        work_dir = _new_workdir()
        _extract_zip_to(folder_zip, work_dir)
        result = batch_rename_by_excel(work_dir, excel_path, old_col, new_col)
        os.remove(excel_path)
        zip_name = _zip_workdir(work_dir, "renamed")
        _clean_logs(result, work_dir)
        return ok(data={**result, "download_url": _download_url(zip_name)})
    except Exception as e:
        return fail(str(e))


# ======================== 模块4：Excel 专项 ========================

@bp.route("/api/mod4/env-check", methods=["GET"])
@module_exception_guard("office_tools")
@login_required
def api_mod4_env_check():
    """检查文档转换引擎环境（Office COM / LibreOffice 任一可用即可转 PDF）"""
    st = engine_status()
    return ok(data={
        "has_win32com": bool(st.get("ms_office")),
        "has_libreoffice": bool(st.get("libreoffice")),
        "engine": st.get("engine"),
        "message": st.get("message"),
    })


@bp.route("/api/mod4/replace-cell", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_mod4_replace_cell():
    """单元格替换：上传文件夹 zip"""
    try:
        folder_zip = _require_zip()
        cell_addr = (request.form.get("cell_addr") or "A1").strip().upper()
        old_text = (request.form.get("old_text") or "").strip()
        new_text = request.form.get("new_text") or ""
        if not old_text:
            return fail("查找文字不能为空")
        work_dir = _new_workdir()
        _extract_zip_to(folder_zip, work_dir)
        result = replace_cell_content(work_dir, cell_addr, old_text, new_text)
        zip_name = _zip_workdir(work_dir, "cells")
        _clean_logs(result, work_dir)
        return ok(data={**result, "download_url": _download_url(zip_name)})
    except RuntimeError as e:
        return fail(str(e), code=503)
    except Exception as e:
        return fail(str(e))


@bp.route("/api/mod4/to-pdf", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_mod4_to_pdf():
    """Excel转PDF：上传文件夹 zip（需服务器装有 Microsoft Office 或 LibreOffice）"""
    try:
        folder_zip = _require_zip()
        work_dir = _new_workdir()
        _extract_zip_to(folder_zip, work_dir)
        result = batch_excel_to_pdf(work_dir)
        zip_name = _zip_workdir(work_dir, "pdf")
        _clean_logs(result, work_dir)
        return ok(data={**result, "download_url": _download_url(zip_name)})
    except RuntimeError as e:
        return fail(str(e), code=503)
    except Exception as e:
        return fail(str(e))


@bp.route("/api/mod4/delete-sheets", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_mod4_delete_sheets():
    """批量删除分表：上传包含Excel的文件夹 zip"""
    try:
        mode = request.form.get("mode", "folder").strip()
        sheet_names_raw = (request.form.get("sheet_names") or "").strip()
        sheet_names = [n.strip() for n in sheet_names_raw.split(",") if n.strip()]
        if not sheet_names:
            return fail("请输入待删除的分表名称")
        folder_zip = _require_zip()
        work_dir = _new_workdir()
        _extract_zip_to(folder_zip, work_dir)
        file_list = [os.path.join(work_dir, item) for item in os.listdir(work_dir)
                     if item.lower().endswith(EXCEL_EXTENSIONS)]
        if not file_list:
            return fail("压缩包内未找到Excel文件")
        result = delete_sheets(file_list, sheet_names)
        zip_name = _zip_workdir(work_dir, "sheets_del")
        _clean_logs(result, work_dir)
        return ok(data={**result, "download_url": _download_url(zip_name)})
    except Exception as e:
        return fail(str(e))


# ======================== 模块5：文件搬运 ========================

@bp.route("/api/mod5/movecopy", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_mod5_movecopy():
    """批量移动/复制：上传源文件夹 zip + 映射Excel"""
    try:
        folder_zip = _require_zip()
        if "excel" not in request.files:
            return fail("请上传映射Excel文件")
        action = request.form.get("action", "copy").strip()
        recursive = request.form.get("recursive", "true").lower() == "true"
        excel_path = save_upload_file(request.files["excel"])
        work_dir = _new_workdir()
        _extract_zip_to(folder_zip, work_dir)
        target_dir = work_dir  # 整理结果直接放在工作根下，随 zip 一并下载
        result = batch_move_copy(work_dir, target_dir, excel_path, action, recursive)
        os.remove(excel_path)
        zip_name = _zip_workdir(work_dir, "organized")
        _clean_logs(result, work_dir)
        return ok(data={**result, "download_url": _download_url(zip_name)})
    except Exception as e:
        return fail(str(e))


@bp.route("/api/mod5/extract-paths", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_mod5_extract_paths():
    """提取文件路径：上传文件夹 zip"""
    try:
        folder_zip = _require_zip()
        path_type = request.form.get("path_type", "absolute").strip()
        recursive = request.form.get("recursive", "true").lower() == "true"
        work_dir = _new_workdir()
        _extract_zip_to(folder_zip, work_dir)
        result = extract_file_paths(work_dir, path_type, recursive)
        token = uuid.uuid4().hex
        _RESULT_CACHE[token] = {"items": result, "root_dir": work_dir, "ts": time.time()}
        return ok(data={"items": result, "count": len(result), "token": token})
    except Exception as e:
        return fail(str(e))


@bp.route("/api/mod5/export-paths", methods=["POST"])
@module_exception_guard("office_tools")
@login_required
def api_mod5_export_paths():
    """导出路径清单：用 token 回查"""
    try:
        token = request.form.get("token") or (request.get_json(silent=True) or {}).get("token")
        cached = _RESULT_CACHE.get(token)
        if not cached:
            return fail("提取结果已失效，请重新执行提取")
        items = cached["items"]
        file_type = request.form.get("file_type") or (request.get_json(silent=True) or {}).get("file_type") or "xlsx"
        if not items:
            return fail("请先执行提取")
        headers = ["文件名", "完整路径", "文件类型"]
        rows = [[r["filename"], r["path"], r["type"]] for r in items]
        fname = _build_table_file(headers, rows, file_type)
        return ok(data={"download_url": f"/office_tools/api/download?file={fname}"})
    except Exception as e:
        return fail(str(e))


# ======================== 模板下载 ========================

@bp.route("/api/template/<template_type>")
@module_exception_guard("office_tools")
@login_required
def api_template(template_type):
    """下载模板"""
    #template_type 原先完全不校验 —— 未知类型会静默走 service 的 else
    # 分支生成一张空表并"成功"返回；更糟的是它被直接拼进文件名
    # （template_{template_type}.xlsx），构造 "../../xxx" 可写出任意路径。
    # 这里白名单收紧，未登记的类型直接 400。
    if template_type not in _TEMPLATE_TYPES:
        return fail("未知的模板类型：%s" % template_type, code=400)
    try:
        wb = generate_template(template_type)
        filepath = os.path.join(UPLOAD_FOLDER, f"template_{template_type}.xlsx")
        wb.save(filepath)
        wb.close()
        return send_file(filepath, as_attachment=True, download_name=f"template_{template_type}.xlsx")
    except Exception as e:
        return fail(str(e))


# ======================== 文件下载 ========================

@bp.route("/api/download")
@module_exception_guard("office_tools")
@login_required
def api_download():
    """下载生成的文件（含结果 zip）"""
    filename = request.args.get("file", "")
    if not filename or ".." in filename or "/" in filename or "\\" in filename:
        return fail("无效的文件名")
    filepath = os.path.join(UPLOAD_FOLDER, filename)
    if not os.path.exists(filepath):
        return fail("文件不存在或已过期")
    return send_file(filepath, as_attachment=True)
