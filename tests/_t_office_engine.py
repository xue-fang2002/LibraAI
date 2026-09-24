"""回归：Office 转换引擎双后端（Office COM / LibreOffice headless）+ 免 Office 的单元格替换。

覆盖：
1. engine_status() 探测结果合理
2. replace_cell_content 走 openpyxl（不依赖任何办公软件）
3. convert_to_pdf 走真实 Office COM（本机装了 Office）
4. convert_to_pdf 在「无 Office」时自动降级 LibreOffice（用桩进程模拟 soffice）
5. 两个引擎都不可用时抛 EngineUnavailable（Runtime 子类 -> 路由会转 503）
6. toolbox office_to_pdf 集成返回 ok

用法：D:/anaconda/python.exe _t_office_engine.py
"""
import os
import sys
import subprocess
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from core import office_engine as oe

RESULTS = []


def check(name, cond, extra=""):
    RESULTS.append((name, bool(cond)))
    print("[%s] %s %s" % ("PASS" if cond else "FAIL", name, extra or ""))


def make_xlsx(path, text="上海分公司"):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = text
    ws["A2"] = None          # 空单元格，用于验证 skip
    wb.save(path)
    return path


def main():
    tmp = tempfile.mkdtemp(prefix="t_office_engine_")

    # ---------- 1. 引擎探测 ----------
    st = oe.engine_status()
    check("engine_status 返回 engine 字段", "engine" in st)
    check("engine_status 返回 message 字段", bool(st.get("message")))
    check("engine_status 同时报告两种引擎",
          "has_win32com" in st and "has_libreoffice" in st)
    print("    探测结果: engine=%s, ms_office=%s, libreoffice=%s"
          % (st.get("engine"), st.get("ms_office"), st.get("libreoffice")))

    # ---------- 2. 单元格替换走 openpyxl（免 Office） ----------
    from modules.office_tools.service import replace_cell_content

    d = os.path.join(tmp, "cells")
    os.makedirs(d, exist_ok=True)
    xlsx = make_xlsx(os.path.join(d, "t.xlsx"))

    r = replace_cell_content(d, "A1", "上海", "北京")
    from openpyxl import load_workbook
    val = load_workbook(xlsx).active["A1"].value
    check("openpyxl 替换生效", val == "北京分公司", "实际=%r" % val)
    check("openpyxl 替换返回 success",
          any(x["status"] == "success" for x in r["log"]))

    r2 = replace_cell_content(d, "A1", "不存在的文字", "X")
    check("无匹配时返回 skip",
          any(x["status"] == "skip" and "无匹配" in x["msg"] for x in r2["log"]))

    r3 = replace_cell_content(d, "A2", "任意", "X")
    check("空单元格返回 skip",
          any(x["status"] == "skip" and "为空" in x["msg"] for x in r3["log"]))

    # ---------- 3. 真实 Office COM 转 PDF ----------
    com_ok = bool(st.get("ms_office"))
    if com_ok:
        src = make_xlsx(os.path.join(tmp, "com_src.xlsx"), "COM 测试")
        outdir = os.path.join(tmp, "com_out")
        os.makedirs(outdir, exist_ok=True)
        try:
            pdf = oe.convert_to_pdf(src, outdir)
            check("Office COM 转 PDF 成功",
                  os.path.isfile(pdf) and pdf.lower().endswith(".pdf"), pdf)
        except Exception as e:
            check("Office COM 转 PDF 成功", False, str(e)[:200])
    else:
        print("[SKIP] 本机无 Office COM，跳过真实转换用例")

    # ---------- 4. 无 Office 时降级 LibreOffice ----------
    src2 = make_xlsx(os.path.join(tmp, "lo_src.xlsx"), "LO 测试")
    outdir2 = os.path.join(tmp, "lo_out")
    os.makedirs(outdir2, exist_ok=True)

    calls = []
    orig_run = oe.subprocess.run
    orig_registered = oe._office_registered
    orig_cache = oe._soffice_cache

    def fake_run(cmd, **kw):
        calls.append(cmd)
        outdir = cmd[cmd.index("--outdir") + 1]
        src = cmd[-1]
        out = os.path.join(
            outdir, os.path.splitext(os.path.basename(src))[0] + ".pdf")
        with open(out, "wb") as f:
            f.write(b"%PDF-1.4 fake")
        return subprocess.CompletedProcess(cmd, 0, b"", b"")

    try:
        oe._office_registered = lambda prog_id: False   # 模拟未装 Office
        oe._soffice_cache = "/fake/soffice"
        oe.subprocess.run = fake_run

        pdf2 = oe.convert_to_pdf(src2, outdir2)
        check("无 Office 时降级 LibreOffice 成功",
              os.path.isfile(pdf2) and pdf2.lower().endswith(".pdf"), pdf2)
        check("LibreOffice 命令行含 --convert-to pdf",
              calls and "--convert-to" in calls[0] and "pdf" in calls[0])
        check("LibreOffice 命令行含 --outdir",
              calls and "--outdir" in calls[0])
        check("LibreOffice 带独立 UserInstallation（避免 profile 抢锁）",
              calls and any("--env:UserInstallation" in c[:6] or
                            c.startswith("-env:UserInstallation") for c in calls[0]))
    except Exception as e:
        check("无 Office 时降级 LibreOffice 成功", False, str(e)[:200])
    finally:
        oe.subprocess.run = orig_run
        oe._office_registered = orig_registered
        oe._soffice_cache = orig_cache

    # ---------- 5. 两个引擎都不可用 ----------
    try:
        oe._office_registered = lambda prog_id: False
        oe._soffice_cache = ""
        try:
            oe.convert_to_pdf(src2, outdir2)
            check("无引擎时抛 EngineUnavailable", False, "竟然没抛异常")
        except oe.EngineUnavailable as e:
            check("无引擎时抛 EngineUnavailable", True)
            check("EngineUnavailable 是 RuntimeError 子类（路由会转 503）",
                  isinstance(e, RuntimeError))
            check("无引擎提示含安装指引",
                  "libreoffice" in str(e).lower() or "Office" in str(e))
    finally:
        oe._office_registered = orig_registered
        oe._soffice_cache = orig_cache

    # ---------- 6. toolbox office_to_pdf 集成 ----------
    try:
        import app as _app_mod
        _r = _app_mod.create_app()
        app = _r[0] if isinstance(_r, tuple) else _r
        with app.app_context():
            from modules.toolbox.tools._common import office_to_pdf
            d3 = os.path.join(tmp, "tb")
            os.makedirs(d3, exist_ok=True)
            src3 = make_xlsx(os.path.join(d3, "tb.xlsx"), "TB")
            resp = office_to_pdf([src3], d3, None, ".pdf")
            data = resp.get_json()
            check("toolbox office_to_pdf 返回 ok",
                  data.get("code") in (0, 200) or data.get("ok") is True
                  or "results" in (data.get("data") or {}),
                  str(data)[:160])
    except Exception as e:
        check("toolbox office_to_pdf 返回 ok", False, str(e)[:200])

    # ---------- 汇总 ----------
    total = len(RESULTS)
    passed = sum(1 for _, ok_ in RESULTS if ok_)
    print("\n==== %d/%d passed ====" % (passed, total))
    return 0 if passed == total else 1


if __name__ == "__main__":
    sys.exit(main())
